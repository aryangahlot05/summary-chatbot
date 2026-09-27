import base64
from pathlib import Path
from typing import Tuple, Optional
import pypdf
import docx
from PIL import Image

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
DOC_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt"}
ALLOWED_EXTENSIONS = IMAGE_EXTENSIONS | DOC_EXTENSIONS

def is_image_file(filename: str) -> bool:
    ext = Path(filename).suffix.lower()
    return ext in IMAGE_EXTENSIONS

def is_doc_file(filename: str) -> bool:
    ext = Path(filename).suffix.lower()
    return ext in DOC_EXTENSIONS

def parse_file(file_path: Path) -> Tuple[str, Optional[str], dict]:
    """
    Parses an uploaded file.
    Returns:
        tuple (extracted_text, base64_image_data_uri_or_none, metadata)
    """
    ext = file_path.suffix.lower()
    metadata = {
        "filename": file_path.name,
        "extension": ext,
        "size_bytes": file_path.stat().st_size,
        "is_image": ext in IMAGE_EXTENSIONS
    }

    if ext in IMAGE_EXTENSIONS:
        return _parse_image(file_path, metadata)
    elif ext == ".pdf":
        return _parse_pdf(file_path, metadata)
    elif ext == ".docx":
        return _parse_docx(file_path, metadata)
    elif ext == ".doc":
        return _parse_doc(file_path, metadata)
    elif ext == ".txt":
        return _parse_txt(file_path, metadata)
    else:
        raise ValueError(f"Unsupported file format: {ext}")

def _parse_image(file_path: Path, metadata: dict) -> Tuple[str, str, dict]:
    with Image.open(file_path) as img:
        width, height = img.size
        format_name = img.format.lower() if img.format else "jpeg"
        metadata["dimensions"] = f"{width}x{height}"
        metadata["format"] = format_name

    with open(file_path, "rb") as f:
        data = f.read()
        b64 = base64.b64encode(data).decode("utf-8")

    mime = "image/png"
    if file_path.suffix.lower() in [".jpg", ".jpeg"]:
        mime = "image/jpeg"
    elif file_path.suffix.lower() == ".webp":
        mime = "image/webp"

    data_uri = f"data:{mime};base64,{b64}"
    summary_text = f"[Image File: {file_path.name}, Size: {metadata['dimensions']}, Format: {metadata.get('format', 'image')}]"
    return summary_text, data_uri, metadata

def _parse_pdf(file_path: Path, metadata: dict) -> Tuple[str, None, dict]:
    reader = pypdf.PdfReader(str(file_path))
    num_pages = len(reader.pages)
    metadata["page_count"] = num_pages

    extracted_pages = []
    for idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            extracted_pages.append(f"--- Page {idx + 1} ---\n{text.strip()}")

    full_text = "\n\n".join(extracted_pages)
    if not full_text.strip():
        full_text = f"[Note: PDF {file_path.name} has {num_pages} pages, but contains no extractable text layer (may be scanned).]"

    return full_text, None, metadata

def _parse_docx(file_path: Path, metadata: dict) -> Tuple[str, None, dict]:
    doc = docx.Document(str(file_path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also parse tables if present
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            table_texts.append(" | ".join(row_cells))

    all_content = []
    if paragraphs:
        all_content.append("\n\n".join(paragraphs))
    if table_texts:
        all_content.append("\n--- Tables ---\n" + "\n".join(table_texts))

    full_text = "\n\n".join(all_content)
    metadata["paragraph_count"] = len(paragraphs)
    metadata["table_count"] = len(doc.tables)

    return full_text, None, metadata

def _parse_doc(file_path: Path, metadata: dict) -> Tuple[str, None, dict]:
    # Legacy .doc binary parsing fallback
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
    except Exception:
        with open(file_path, "r", encoding="latin-1", errors="ignore") as f:
            content = f.read()

    # Filter readable characters
    readable_chars = "".join([c for c in content if c.isprintable() or c in "\n\r\t"])
    lines = [line.strip() for line in readable_chars.splitlines() if len(line.strip()) > 3]
    text = "\n".join(lines)
    if not text.strip():
        text = f"[Legacy .doc file: {file_path.name}. For optimal parsing, please convert to .docx format.]"

    return text, None, metadata

def _parse_txt(file_path: Path, metadata: dict) -> Tuple[str, None, dict]:
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1", errors="replace") as f:
            text = f.read()

    metadata["character_count"] = len(text)
    return text, None, metadata
